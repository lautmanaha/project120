#!/usr/bin/env python3
"""פרויקט 120 - ייצוא מסד הסקרים לקובץ Excel קריא (exports/project120_polls.xlsx)."""
from __future__ import annotations

import sqlite3
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DB = ROOT / "db" / "polls.sqlite"
OUT = ROOT / "exports" / "project120_polls.xlsx"
sys.path.insert(0, str(ROOT / "db"))
from build_db import PARTIES  # noqa: E402

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(bold=True, color="FFFFFF", name="Arial")
BODY_FONT = Font(name="Arial")


def sheet(wb: Workbook, title: str, header: list[str], rows: list, widths: dict | None = None):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BODY_FONT
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(10, min(28, len(str(h)) + 4)))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def main() -> int:
    conn = sqlite3.connect(DB)
    wb = Workbook()
    wb.remove(wb.active)

    # 1. סקרים - שורה לכל סקר עם מנדטים בשאלה הראשית
    polls = conn.execute("""
        SELECT p.reference_number, ps.name, pb.name, pb.type, p.fieldwork_start, p.fieldwork_end,
               m.publish_date, p.method, p.initial_sample, p.respondents, p.response_rate_pct,
               p.margin_error_pct, q.undecided_pct, p.in_model, p.exclude_reason, q.question_id
        FROM polls p
        LEFT JOIN pollsters ps ON ps.pollster_id=p.pollster_id
        LEFT JOIN publishers pb ON pb.publisher_id=p.publisher_id
        LEFT JOIN polls_meta m ON m.reference_number=p.reference_number
        LEFT JOIN poll_questions q ON q.reference_number=p.reference_number AND q.type='main'
        ORDER BY p.fieldwork_end, p.reference_number""").fetchall()
    method_he = {"internet": "אינטרנט", "phone": "טלפון", "mixed": "משולב", "unknown": "לא צוין"}
    rows = []
    for r in polls:
        seats = dict(conn.execute("SELECT party, seats FROM poll_results WHERE question_id=?", (r[15],)).fetchall()) if r[15] else {}
        rows.append([r[0], r[1], r[2], r[3], r[4], r[5], r[6], method_he.get(r[7], r[7]), r[8], r[9], r[10], r[11], r[12],
                     "כן" if r[13] else "לא", r[14]] + [seats.get(p) for p in PARTIES])
    header = ["מס' אסמכתא", "מכון", "ערוץ פרסום", "סוג ערוץ", "תחילת שדה", "סיום שדה", "פורסם באתר הוועדה",
              "שיטה", "מדגם התחלתי", "משיבים", "% היענות", "טעות דגימה", "% לא החליטו", "במודל", "סיבת אי-הכללה"] + PARTIES
    sheet(wb, "סקרים - מנדטים", header, rows, {"מכון": 22, "ערוץ פרסום": 20, "סיבת אי-הכללה": 30})

    # 2. אחוזים
    rows = []
    for r in polls:
        pct = dict(conn.execute("SELECT party, pct FROM poll_results WHERE question_id=?", (r[15],)).fetchall()) if r[15] else {}
        rows.append([r[0], r[1], r[2], r[5], r[9]] + [pct.get(p) for p in PARTIES])
    sheet(wb, "סקרים - אחוזים", ["מס' אסמכתא", "מכון", "ערוץ פרסום", "סיום שדה", "משיבים"] + PARTIES, rows,
          {"מכון": 22, "ערוץ פרסום": 20})

    # 3. כל השאלות (כולל תרחישים) - פורמט ארוך
    rows = conn.execute("""
        SELECT r.reference_number, ps.name, pb.name, p.fieldwork_end, q.qid,
               CASE q.type WHEN 'main' THEN 'ראשית' ELSE 'תרחיש' END, q.description,
               r.party, r.party_as_written, r.seats, r.pct, q.undecided_pct, q.notes
        FROM poll_results r JOIN poll_questions q ON q.question_id=r.question_id
        JOIN polls p ON p.reference_number=r.reference_number
        LEFT JOIN pollsters ps ON ps.pollster_id=p.pollster_id
        LEFT JOIN publishers pb ON pb.publisher_id=p.publisher_id
        ORDER BY p.fieldwork_end, r.reference_number, q.qid, r.seats DESC""").fetchall()
    sheet(wb, "כל התוצאות (ארוך)",
          ["מס' אסמכתא", "מכון", "ערוץ", "סיום שדה", "מס' שאלה", "סוג", "נוסח/תרחיש", "מפלגה (קנוני)",
           "מפלגה כפי שנכתב", "מנדטים", "%", "% לא החליטו", "הערות"], rows,
          {"נוסח/תרחיש": 45, "מפלגה כפי שנכתב": 35, "הערות": 40, "מכון": 22})

    # 4. מכונים / ערוצים
    sheet(wb, "מכונים", ["מכון (קנוני)", "מס' סקרים", "כתיבים שנתקלנו בהם"],
          conn.execute("""SELECT name, (SELECT COUNT(*) FROM polls WHERE pollster_id=pollsters.pollster_id), aliases
                          FROM pollsters ORDER BY 2 DESC""").fetchall(), {"כתיבים שנתקלנו בהם": 90, "מכון (קנוני)": 24})
    sheet(wb, "ערוצי פרסום", ["ערוץ (קנוני)", "סוג", "מס' סקרים", "כתיבים"],
          conn.execute("""SELECT name, type, (SELECT COUNT(*) FROM polls WHERE publisher_id=publishers.publisher_id), aliases
                          FROM publishers ORDER BY 3 DESC""").fetchall(), {"כתיבים": 70, "ערוץ (קנוני)": 24})

    # 5. רשומות הוועדה כפי שהן
    sheet(wb, "רשומות הוועדה (גולמי)",
          ["מס' אסמכתא", "עורך הסקר", "מזמין הסקר", "מועד ביצוע", "מועד העברה", "מועד פרסום", "מועד עדכון", "הערות", "קובץ", "קישור"],
          conn.execute("""SELECT reference_number, survey_editor, survey_publisher, survey_date, transfer_date,
                          publish_date, update_date, notes, file_name, file_url FROM polls_meta
                          ORDER BY publish_date, reference_number""").fetchall(),
          {"עורך הסקר": 30, "מזמין הסקר": 30, "קובץ": 36, "קישור": 60})

    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print("נשמר", OUT)
    return 0


if __name__ == "__main__":
    sys.exit(main())
